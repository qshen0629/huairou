#!/usr/bin/env python3
"""生成可编辑 PPTX、同版式 PDF、逐页讲稿与缩略图；不包含假装生成的真人媒体。"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from functools import lru_cache
import math
import os

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "课件"
W, H = 1600, 900
BG, PANEL, WHITE = "101D32", "1C2D46", "F1F4F7"
MUTED, LIME, CYAN, CORAL = "B6C5D8", "DEF97D", "80DCE3", "FFB39E"
BLUE, DARK = "314862", "101D32"
FONT_PATH = next(p for p in [
    os.environ.get("CLASSROOM_FONT", ""),
    str(ROOT / "assets/fonts/NotoSansSC-Regular.ttf"),
    "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
    "/System/Library/Fonts/PingFang.ttc",
    "/System/Library/Fonts/STHeiti Light.ttc",
] if p and Path(p).is_file())


@lru_cache(maxsize=None)
def font(size):
    return ImageFont.truetype(FONT_PATH, int(size))


def wrapped(text, width, size):
    f = font(size)
    result = []
    for paragraph in text.split("\n"):
        current = ""
        for char in paragraph:
            if current and f.getlength(current + char) > width:
                result.append(current)
                current = char
            else:
                current += char
        result.append(current)
    return result


@dataclass
class Slide:
    title: str
    minutes: int
    notes: str
    elements: list = field(default_factory=list)

    def box(self, x, y, w, h, fill=PANEL, radius=24):
        self.elements.append(dict(kind="box", x=x, y=y, w=w, h=h, fill=fill, radius=radius))

    def oval(self, x, y, w, h, fill):
        self.elements.append(dict(kind="oval", x=x, y=y, w=w, h=h, fill=fill))

    def line(self, x1, y1, x2, y2, color=BLUE, weight=3):
        self.elements.append(dict(kind="line", x=x1, y=y1, x2=x2, y2=y2, color=color, weight=weight))

    def text(self, x, y, w, h, text, size=42, color=WHITE, bold=False, align="left", leading=1.28):
        lines = wrapped(text, w, size)
        if len(lines) * size * leading > h + 3:
            raise ValueError(f"文本框高度不足：{self.title} / {text}")
        self.elements.append(dict(kind="text", x=x, y=y, w=w, h=h, text="\n".join(lines),
                                  size=size, color=color, bold=bold, align=align, leading=leading))


SLIDES = []


def new(title, minutes, notes, kicker="科技小侦探"):
    s = Slide(title, minutes, notes)
    SLIDES.append(s)
    s.text(88, 50, 1424, 44, kicker, 24, LIME)
    s.text(88, 114, 1424, 110, title, 64, WHITE, True)
    return s


def bar(s, message, color=LIME, y=728, size=40):
    s.box(88, y, 1424, 86, color, 20)
    s.text(116, y + 16, 1368, 60, message, size, DARK, True)


def cards(s, items, y=276, h=400):
    for i, (tag, title, body, accent) in enumerate(items):
        x = 88 + i * 484
        s.box(x, y, 456, h)
        if h < 350:
            s.text(x + 30, y + 18, 396, 55, tag, 36, accent, True)
            s.text(x + 30, y + 86, 396, 114, title, 39, WHITE, True)
            s.text(x + 30, y + 226, 396, 58, body, 25, MUTED)
        else:
            s.text(x + 30, y + 26, 396, 65, tag, 44, accent, True)
            s.text(x + 30, y + 116, 396, 128, title, 46, WHITE, True)
            s.text(x + 30, y + 272, 396, h - 284, body, 31, MUTED)


def face(s, cx, cy, scale=1, mouth=0, color=CYAN):
    s.box(cx - 97 * scale, cy + 90 * scale, 194 * scale, 110 * scale, color, 34 * scale)
    s.oval(cx - 76 * scale, cy - 105 * scale, 152 * scale, 184 * scale, WHITE)
    s.oval(cx - 40 * scale, cy - 37 * scale, 14 * scale, 14 * scale, DARK)
    s.oval(cx + 25 * scale, cy - 37 * scale, 14 * scale, 14 * scale, DARK)
    if mouth:
        s.oval(cx - 22 * scale, cy + 22 * scale, 44 * scale, (10 + mouth * 24) * scale, DARK)
    else:
        s.line(cx - 21 * scale, cy + 35 * scale, cx + 21 * scale, cy + 35 * scale, DARK, 5 * scale)


def wave(s, x, y, width=500, height=90, accent=CYAN):
    for i in range(45):
        amplitude = (0.2 + 0.8 * abs(math.sin(i * 1.73)) * abs(math.sin(i * 0.16 + 0.2))) * height
        s.box(x + i * width / 45, y + (height - amplitude) / 2, width / 70, amplitude, accent, 3)


def media(s, x, y, w, h, title, detail, accent=CYAN):
    s.box(x, y, w, h)
    s.box(x + 24, y + 24, 10, h - 48, accent, 4)
    s.text(x + 56, y + 34, w - 82, 110, title, 44, WHITE, True)
    s.text(x + 56, y + h - 94, w - 82, 76, detail, 27, MUTED)


def compose():
    s = new("谁在说话？", 1,
        "开场：大家好，今天我们一起拆一个科技魔术——让我的声音和照片变成一个数字分身。接下来会出现合成片段，我们猜完就揭晓。\n"
        "提问：屏幕里的人长得像我、声音也像我，就一定是我本人说的吗？先不公布答案。\n"
        "课堂目标：能说清声音合成与人像动画的基本输入输出；能接受证据不足；能给出一种独立核实方式。")
    s.text(88, 258, 850, 240, "当声音和脸\n都可以被模仿", 84, WHITE, True)
    s.text(88, 550, 820, 140, "一场关于 AI、数字分身\n与证据的科技侦探课", 42, MUTED)
    s.box(1030, 250, 480, 470)
    face(s, 1270, 420, 1.15, mouth=1)
    for i in range(7):
        s.box(1372 + (i % 2) * 28, 333 + i * 40, 25, 25, LIME if i % 2 else CYAN, 2)
    s.text(1080, 676, 380, 50, "相似，是证明吗？", 34, LIME, align="center")
    s.text(88, 754, 900, 46, "六年级 · 40 分钟 · 本课含明确揭晓的合成演示", 26, CYAN)

    s = new("今天，我们来当科技小侦探", 1,
        "请学生记住三项任务：先猜，再拆解，最后核实。猜错很正常，可以回答‘目前还不知道’。\n"
        "提醒：我们展示老师自己的素材，不收集同学的照片和声音。\n"
        "串场：第一关先关掉画面，只用耳朵。")
    cards(s, [("任务 01", "听一听\n谁在说话", "选择 A、B，或不确定", CYAN),
              ("任务 02", "拆一拆\n分身怎么做", "找出需要哪些输入", LIME),
              ("任务 03", "想一想\n怎样去核实", "给出一个能执行的行动", CORAL)])
    bar(s, "侦探可以说：我还需要更多证据。")

    s = new("只听声音，你选哪一段？", 3,
        "课前：在 A 框插入 teacher_real.wav，在 B 框插入 01_voice_ai.wav，二者台词相同，音量接近。此页模板尚未嵌入音频。\n"
        "播放顺序：A、B，各 5–10 秒，必要时重放一轮。全班举手选择：A 是原声 / B 是原声 / 证据不足。\n"
        "邀请两位学生解释依据，比如停顿、音色、呼吸。不要评价谁聪明或谁被骗。\n"
        "注意：有模型合成标识不是身份认证。这轮已经告知包含合成，下一页立即揭晓。")
    for i, label in enumerate(["声音 A", "声音 B"]):
        x = 88 + i * 726
        s.box(x, 280, 698, 390)
        s.text(x + 40, 316, 618, 90, label, 68, CYAN if i == 0 else LIME, True)
        wave(s, x + 48, 452, 600, 100, CYAN if i == 0 else LIME)
        s.text(x + 40, 598, 618, 54, "插入音频 · 听相同的一句话", 29, MUTED)
    bar(s, "A、B，还是证据不足？请说出你的理由。", CYAN, size=39)

    s = new("揭晓：相似的声音可以生成", 2,
        "答案固定为 A 原声、B 合成。若你在上一页交换了媒体顺序，务必同步更改本页答案。\n"
        "讲解：B 的台词不是从我的原声测试录音里剪出来的。模型参考一小段声音和对应文字，再根据新台词生成新的声音。\n"
        "如果学生猜中，追问：这只能说明你发现了这次生成的一些线索。如果下一次更自然呢？\n"
        "可选现场操作：预先加载模型，输入一句学生建议的无害短台词。等待超过约 20 秒就播放课前成片，不把 20 秒说成生成速度保证。")
    s.box(88, 278, 698, 250)
    s.box(814, 278, 698, 250)
    s.text(130, 316, 620, 90, "A：原声录音", 58, CYAN, True)
    s.text(856, 316, 620, 90, "B：AI 合成", 58, LIME, True)
    s.text(130, 430, 610, 70, "老师读出这段话", 37, WHITE)
    s.text(856, 430, 610, 70, "模型生成这段话", 37, WHITE)
    s.text(100, 590, 1400, 94, "参考声音 + 新台词 → 相似音色的新声音", 47, WHITE, True, "center")
    bar(s, "声音像老师 ≠ 这句话由老师本人说过")

    s = new("第二关：一张照片，开口了", 3,
        "课前：左框换成你自己的 teacher.jpg，右框嵌入 02_talking_ai.mp4。本模板只有占位框。\n"
        "先展示静态照片，问：只有一张照片，能拍成刚才这样的说话视频吗？播放 5–10 秒成片，保留持续合成标识。\n"
        "请学生猜输入。揭晓：一张老师的照片 + 合成好的音频，由 SadTalker 生成说话动作与逐帧画面。\n"
        "解释：这是声音驱动的人像动画。它并没有证明我在那个时间、那个地点说过这句话。\n"
        "Mac 现场不等待视频生成。")
    media(s, 88, 274, 570, 420, "一张照片", "替换成老师本人的正面照", CYAN)
    face(s, 372, 445, 0.7, color=CYAN)
    s.text(683, 420, 110, 140, "→", 90, LIME, align="center")
    media(s, 814, 274, 698, 420, "会说话的照片", "插入预生成视频 · 保留 AI 合成标识", LIME)
    s.text(872, 455, 580, 90, "照片 + 音频", 56, LIME, True, "center")
    bar(s, "画面里的“我”，不一定是真实录下来的我。", size=38)

    s = new("第三关：你点头，照片也点头", 2,
        "扩展示范：播放 LivePortrait 的 03_motion_ai.mp4，对照输入照片、驱动动作与生成结果。没有准备此素材，可以跳过本页，把两分钟留给讨论。\n"
        "先问学生：这次是声音在指挥嘴，还是动作视频在指挥表情？\n"
        "说明边界：LivePortrait 把动作/表情迁移到肖像，不会凭 F5-TTS 的新音频自动做匹配口型。不要只把两者混音后当作完成对口型。\n"
        "首帧正面自然表情，动作控制在点头、微笑或眨眼；展示本人已授权素材。")
    for i, (title, detail, color) in enumerate([
        ("输入照片", "老师本人的肖像", CYAN),
        ("驱动动作", "点头、微笑、眨眼", CORAL),
        ("生成结果", "照片跟着动起来", LIME),
    ]):
        x = 88 + 484 * i
        media(s, x, 274, 456, 420, title, detail, color)
        face(s, x + 255, 465, 0.65, mouth=i / 2, color=color)
    bar(s, "这一关由动作视频驱动；不会自动按新声音对口型。", CYAN, size=34)

    s = new("拆开魔术：两条合成流水线", 3,
        "请三位学生依次说出第一条流水线的材料：参考声音、参考逐字稿、新台词。模型输出一段新声音。\n"
        "第二条流水线用老师照片和这段声音，生成会说话的视频。把学生的答案逐步指到图上。\n"
        "重要区分：上课演示的是调用已经训练好的模型做推理。我们不是用十秒钟从头训练了一个懂所有声音的模型。\n"
        "追问：视频里每个组成部分都像真的，整个事件就一定真实发生了吗？")
    for row, (left, center, right, accent) in enumerate([
        ("参考声音、逐字稿\n+ 想说的新台词", "声音合成器", "合成声音", CYAN),
        ("老师照片\n+ 合成声音", "人脸动画器", "说话视频", LIME),
    ]):
        y = 274 + row * 218
        for x, width in [(88, 482), (656, 394), (1136, 376)]:
            s.box(x, y, width, 190)
        s.text(116, y + 35, 430, 132, left, 37, WHITE)
        s.text(685, y + 60, 338, 78, center, 45, accent, True, "center")
        s.text(1160, y + 60, 330, 78, right, 45, accent, True, "center")
        s.text(578, y + 54, 70, 95, "→", 56, MUTED, align="center")
        s.text(1058, y + 54, 70, 95, "→", 56, MUTED, align="center")
    bar(s, "用训练好的模型生成新内容，叫“推理”。", size=39)

    s = new("声音合成器，参考了什么？", 2,
        "可以让学生同一句‘欢迎来到科技课堂’分别用缓慢、轻快的节奏读一遍，体会声音有不同特点。不是让学生把录音交给模型。\n"
        "用简化解释：模型在大量已有数据中学到声音与文字的关系，推理时参考样本中的音色、发音与节奏，再生成新声音。并非每种声学特征都会被准确保留。\n"
        "不要说模型复制了人的大脑、思想或完整身份。此处是直观解释，不讲 Transformer、损失函数或数学细节。")
    cards(s, [("线索 01", "音色", "听起来像谁？", CYAN),
              ("线索 02", "节奏", "快慢、停顿怎样？", LIME),
              ("线索 03", "发音", "字词怎样被说出来？", CORAL)])
    bar(s, "参考声音特点，再生成新的声音。")

    s = new("人脸动画，一帧接着一帧", 2,
        "说明这三张脸是概念示意图，不是模型真实中间输出。让学生观察嘴巴和表情随时间变化。\n"
        "讲解：静态照片给出人物外观，声音或动作信息给出变化线索。模型生成一系列变化的画面，连续播放就形成动画。\n"
        "动画不是一段真实拍摄事件的证明。传统动画也会动，但 AI 能让它更接近真实人像。")
    for i in range(3):
        x = 88 + i * 484
        s.box(x, 274, 456, 420)
        s.text(x + 30, 300, 390, 66, f"第 {i + 1} 帧", 42, CYAN if i == 0 else LIME)
        face(s, x + 228, 480, 0.88, mouth=[0, 0.8, 0.25][i])
    bar(s, "画面会动，不等于拍到了真实发生的事。", size=38)

    s = new("数字分身，就是另一个我吗？", 2,
        "让学生先讨论 30 秒，再找两位回答。\n"
        "引导：它能模仿我的声音与外观，但一段相似的生成视频，不能单凭像我就证明我在场、我同意或者我确实说过。\n"
        "可选思考：AI 合成的科普知识一定是错的吗？真实拍摄的视频里说的话一定正确吗？画面是否合成与事情是否属实，是两个不同问题，都要看相应证据。\n"
        "这不是关于‘AI 有没有灵魂’的辩论；把思考落在证据和授权上。也避免说所有视频都是假的。\n"
        "串场：那我们能不能靠找破绽来一招验真？")
    s.box(88, 280, 698, 398)
    s.box(814, 280, 698, 398)
    s.text(128, 318, 618, 86, "可以模仿", 58, CYAN, True)
    s.text(128, 438, 618, 190, "声音的特点\n脸部的外观和动作", 43, WHITE)
    s.text(854, 318, 618, 86, "不能单凭相似证明", 49, CORAL, True)
    s.text(854, 438, 618, 190, "本人当时在场\n本人说过、本人同意", 43, WHITE)
    bar(s, "它生成画面和声音；我们仍要核实事件。", size=39)

    s = new("找破绽，有用但不够", 3,
        "播放刚才的短片一次，问：你观察到什么？可提示嘴形与声音、脸部边缘、画面是否抖动。不要捏造视频里没有出现的缺陷。\n"
        "提醒：原始视频也可能因压缩、延迟、光照等出现异常；异常只能引起怀疑，不能一项就判定造假。\n"
        "反过来，没有看见异常也不能判真。眨眼、指定动作、口令或检测器分数都不是绝对验真法。\n"
        "课堂不用虚构检测准确率，也不拿无标注小样本演示‘检测器百分之百识别’。")
    cards(s, [("可以观察", "嘴形与声音", "有没有明显不同步？", CYAN),
              ("可以观察", "边缘与细节", "有没有异常变形？", CYAN),
              ("也要知道", "真实视频\n也可能异常", "压缩、延迟都可能影响画面", CORAL)])
    bar(s, "有异常：值得核实。没异常：仍需要证据。", CORAL, size=38)

    s = new("侦探任务：接下来该做什么？", 4,
        "明确这是虚构情境，不是发生在某个同学家的真实案件。解释验证码是手机收到的一次性安全密码。\n"
        "情境：一个视频里的脸和声音都很像熟悉的家人，对方要求你发手机验证码，并叫你不要告诉别人。\n"
        "学生两三人讨论一分钟，选择 A/B/C 并说理由。不给他们观看实际冒充家人的合成视频。\n"
        "追问：如果你还不能确定视频真伪，能不能先决定不发验证码？可以。\n"
        "关键是安全行动和独立核实，不是考学生看破每个假视频。")
    s.box(88, 258, 1424, 180, PANEL)
    s.text(120, 280, 1350, 50, "虚构情境 · 熟悉的脸，熟悉的声音", 29, CORAL)
    s.text(120, 342, 1350, 80, "“把手机验证码发给我，先别告诉别人。”", 49, WHITE, True)
    for i, (tag, text, accent) in enumerate([
        ("A", "看着像家人，\n立刻发给他", CORAL),
        ("B", "多看几遍嘴巴，\n没有破绽就发", CYAN),
        ("C", "先不发，找可信任的大人，\n用已知联系方式核实", LIME),
    ]):
        x = 88 + i * 484
        s.box(x, 472, 456, 238)
        s.text(x + 28, 494, 396, 76, tag, 58, accent, True)
        s.text(x + 28, 588, 400, 112, text, 30, WHITE)
    bar(s, "现在不必急着判真假，先决定下一步。", size=40)

    s = new("证据不够，也可以先做对的事", 3,
        "建议选择 C。但最关键的原则是：验证码不转发，即使对方确实是熟人也一样。\n"
        "为什么 A 不够：相似的脸与声音并不能单独认证身份。为什么 B 不够：看不出破绽也不能证明真。\n"
        "如何核实：找身边可信任的大人，用早已保存的联系方式，或当面确认。不要用可疑视频刚提供的新号码，也不要只在同一个可疑聊天里问‘你是真的吗’。\n"
        "即使联系不到，也不用立即按对方要求行动。能说‘证据不足’就是有价值的判断。")
    s.text(88, 270, 1424, 114, "选择 C：先不发送验证码", 70, LIME, True)
    cards(s, [("第一步", "暂停操作", "验证码不转发", CORAL),
              ("第二步", "找可信任\n的大人", "一起判断和处理", CYAN),
              ("第三步", "换可信渠道\n核实", "用早已知道的联系方式", LIME)], y=418, h=290)
    bar(s, "视频真伪暂时不明，也不影响先保护自己。", size=38)

    s = new("记住四个字：停、想、问、核", 3,
        "让全班一起读四个字。\n停：先暂停转发、付款、发送验证码等动作。\n想：这条消息让我做什么？是否催促、保密或索要账号信息？\n问：找身边可信任的大人，不独自处理。\n核：用已知联系方式、当面或其他独立可信渠道核实。\n"
        "邀请学生把第 12 页的情境套进四步。区分‘再看一遍’和‘得到新的独立证据’。\n"
        "视频通话里让对方做动作只能增加线索，不应作为绝对保证。")
    for i, (word, detail, accent) in enumerate([
        ("停", "先暂停操作", CORAL), ("想", "他要我做什么", CYAN),
        ("问", "找可信任的大人", CYAN), ("核", "换可信渠道确认", LIME),
    ]):
        x = 88 + i * 363
        s.box(x, 280, 335, 414)
        s.text(x + 24, 316, 287, 210, word, 148, accent, True, "center")
        s.text(x + 24, 580, 287, 98, detail, 33, WHITE, align="center")
    bar(s, "核实用已知渠道，别只跟着可疑消息里的新号码走。", size=34)

    s = new("同一项技术，可以怎样帮助人？", 2,
        "让学生把技术想成工具，提出有益用途。可讨论：协助失声者表达、给知识内容配多语言讲解、制作获授权的角色动画。这些是用途方向，不表示课堂这套模型已通过医疗用途验证。\n"
        "追问：如果把同学做成搞笑视频，即使好笑也可以直接发出去吗？讨论同意、尊重和传播后果。\n"
        "把结论说成可操作原则：用自己或明确授权的素材，清楚标注合成用途。不要把网络上可看到的照片当作别人已经同意。")
    cards(s, [("表达", "让交流\n多一种办法", "怎样帮助需要的人？", CYAN),
              ("知识", "为讲解\n配上多种语言", "怎样让更多人听懂？", LIME),
              ("创作", "制作获授权\n的角色动画", "怎样尊重被模仿的人？", CORAL)])
    bar(s, "先获得同意，再清楚标注合成用途。")

    s = new("离堂挑战：带走三个答案", 3,
        "每题给约 20 秒思考，再请学生回答理由。\n"
        "1. 视频里的人会眨眼，就一定是真的吗？答：不是，动作也能合成。\n"
        "2. 看不出视频破绽，下一步怎么办？答：不能仅凭这点判真，按消息要求的风险先暂停，再向大人求助，用已知可信渠道核实。\n"
        "3. 想把同学做成搞笑分身，可以直接发吗？答：不可以，应先取得明确同意、考虑尊重和传播后果，课堂不收集学生素材。\n"
        "评价目标：能否给出具体核实行动，而不是能否永远猜中真假。")
    cards(s, [("01", "会眨眼，\n就一定真吗？", "为什么？", CYAN),
              ("02", "看不出破绽，\n下一步做什么？", "给出一个具体行动", LIME),
              ("03", "同学的分身，\n可以直接发吗？", "还需要考虑什么？", CORAL)])
    bar(s, "你不必永远猜中，但可以学会怎样核实。", size=39)

    s = new("让科技有趣，让判断有依据", 1,
        "收束：今天我们看见声音可以被生成，照片也能动起来。这些能力很有趣，但像某个人不等于证明那个人真实表达过。\n"
        "请全班再说一次：停、想、问、核。\n"
        "可选家庭小任务：回家教大人一个核实方法。不要要求学生制作同学分身或提交家人的照片、声音。\n"
        "结束，不在课堂展示后面的教师技术附录。")
    s.text(88, 280, 1424, 260, "眼见和耳听，\n都要和证据做朋友。", 92, WHITE, True)
    s.text(88, 605, 1424, 90, "会使用 AI，也会核实信息。", 54, CYAN)
    bar(s, "停  ·  想  ·  问  ·  核", size=46)

    s = new("教师附录：演示制作顺序", 0,
        "本页不计入 40 分钟课堂。详细命令见根目录 README.md。\n"
        "素材：本人 8–10 秒参考语音 + 准确逐字稿；同台词测试原声；正面照；可选 5–8 秒动作视频。\n"
        "顺序：F5-TTS 生成 01_voice_ai.wav；SadTalker 使用照片和该音频生成 02_talking_ai.mp4；可选 LivePortrait 生成 03_motion_ai.mp4。\n"
        "Mac 上视频提前生成。现场仅尝试已在该机器上成功演练的短句语音推理；超过约 20 秒切换课前素材。\n"
        "所有模型首次下载、环境安装与素材嵌入都安排在课前。", kicker="教师附录 · 不投给学生的技术细节")
    cards(s, [("先完成", "F5-TTS\n声音克隆", "参考录音 + 逐字稿 + 台词", CYAN),
              ("主展示", "SadTalker\n照片说话", "照片 + 上一步生成的声音", LIME),
              ("有余力再做", "LivePortrait\n动作迁移", "照片 + 动作视频，非音频驱动", CORAL)])
    bar(s, "稳定的课前短片 + 一次可选的现场短句演示", CYAN, size=38)

    s = new("教师附录：Mac 能力边界", 0,
        "苹果芯片 Mac：F5-TTS 官方有安装路径，代码支持 MPS；LivePortrait 官方有人像模式 Mac 支持，但提示可能明显慢于高端 NVIDIA GPU。\n"
        "SadTalker 当前 inference.py 只判断 CUDA / CPU，本方案在 Mac 明确使用 CPU，不能称为 MPS 加速。\n"
        "未在用户 Mac 上实测。不能承诺实时视频生成、具体耗时或一键安装成功。\n"
        "许可证：F5-TTS 代码 MIT，官方权重 CC-BY-NC-4.0；SadTalker 主项目 Apache-2.0，组件另有条款；LivePortrait 代码 MIT，但随附 InsightFace 模型限非商业研究。教学用途不自动覆盖所有模型授权情形。\n"
        "项目源码已下载，真人素材、权重及实际推理结果尚缺。", kicker="教师附录 · 安装和授权以原项目为准")
    for i, (name, info, accent) in enumerate([
        ("F5-TTS", "苹果芯片 / MPS；先在你的 Mac 实际演练", CYAN),
        ("SadTalker", "Mac 走 CPU；声音驱动；课前生成视频", LIME),
        ("LivePortrait", "苹果芯片人像模式；动作驱动；不自动对音频口型", CORAL),
    ]):
        y = 274 + i * 140
        s.box(88, y, 1424, 116)
        s.text(116, y + 24, 360, 72, name, 44, accent, True)
        s.text(510, y + 32, 970, 66, info, 34, WHITE)
    bar(s, "代码许可与模型许可分开看；本套未做你的 Mac 推理实测。", CORAL, size=33)

    s = new("教师附录：项目与资料来源", 0,
        "来源核对日期：2026-10-14。链接均为项目官方仓库、官方文档或模型卡。\n"
        "F5-TTS：https://github.com/SWivid/F5-TTS；https://huggingface.co/SWivid/F5-TTS\n"
        "SadTalker：https://github.com/OpenTalker/SadTalker；Mac 安装记录：https://github.com/OpenTalker/SadTalker/blob/main/docs/install.md\n"
        "LivePortrait：https://github.com/KlingAIResearch/LivePortrait；许可证：https://github.com/KlingAIResearch/LivePortrait/blob/main/LICENSE\n"
        "备选项目和未优先选用原因见 README.md。源码版本标识见 projects/SOURCES.json。\n"
        "本课的原理图是面向六年级的简化示意，不是模型架构图或真实中间张量可视化。PDF 与缩略图由同一版面数据生成，不是 PowerPoint 渲染截图。", kicker="教师附录 · 官方来源 · 2026-10-14 核对")
    sources = [
        ("F5-TTS · 官方代码与模型卡", "https://github.com/SWivid/F5-TTS\nhttps://huggingface.co/SWivid/F5-TTS"),
        ("SadTalker · 官方代码与 Mac 安装说明", "https://github.com/OpenTalker/SadTalker\nhttps://github.com/OpenTalker/SadTalker/blob/main/docs/install.md"),
        ("LivePortrait · 官方代码与组件许可", "https://github.com/KlingAIResearch/LivePortrait\nhttps://github.com/KlingAIResearch/LivePortrait/blob/main/LICENSE"),
    ]
    for i, (name, link) in enumerate(sources):
        y = 256 + i * 180
        s.box(88, y, 1424, 160)
        s.text(116, y + 14, 1368, 62, name, 36, CYAN if i % 2 == 0 else LIME, True)
        s.text(116, y + 77, 1368, 76, link, 25, WHITE)

    for i, s in enumerate(SLIDES, 1):
        s.line(88, 841, 1512, 841, BLUE, 1.5)
        s.text(88, 856, 1260, 30, "科技小侦探 · 相似是线索，核实需要证据" if s.minutes else
               "教师资料 · 课件中的真人音视频位置尚待填入", 18, MUTED)
        s.text(1368, 850, 144, 38, f"{i:02d} / {len(SLIDES):02d}", 24, MUTED, align="right")


def ppt(slides):
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13 + 1 / 3), Inches(7.5)
    p.core_properties.title = "谁在说话？——六年级 AI 数字分身科技课"
    p.core_properties.subject = "40 分钟 · 声音合成、人像动画与信息核实"
    p.core_properties.author = "科技展示课"
    scale = 1 / 120
    for i, item in enumerate(slides, 1):
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        for e in item.elements:
            x, y = Inches(e["x"] * scale), Inches(e["y"] * scale)
            if e["kind"] in ("box", "oval"):
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if e["kind"] == "oval" else
                    MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE,
                    x, y, Inches(e["w"] * scale), Inches(e["h"] * scale))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
                shape.line.fill.background()
                if e["kind"] == "box" and e["radius"]:
                    shape.adjustments[0] = min(0.5, e["radius"] / min(e["w"], e["h"]))
            elif e["kind"] == "line":
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y,
                    Inches(e["x2"] * scale), Inches(e["y2"] * scale))
                shape.line.color.rgb = RGBColor.from_string(e["color"])
                shape.line.width = Pt(e["weight"] * 0.6)
            else:
                shape = slide.shapes.add_textbox(x, y, Inches(e["w"] * scale), Inches(e["h"] * scale))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.word_wrap = False
                for n, line in enumerate(e["text"].split("\n")):
                    para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                    para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[e["align"]]
                    para.line_spacing = Pt(e["size"] * e["leading"] * 0.6)
                    para.space_before = para.space_after = Pt(0)
                    run = para.add_run()
                    run.text = line
                    run.font.name = "PingFang SC"
                    run.font.size = Pt(e["size"] * 0.6)
                    run.font.bold = e["bold"]
                    run.font.color.rgb = RGBColor.from_string(e["color"])
                    ea = OxmlElement("a:ea")
                    ea.set("typeface", "PingFang SC")
                    run._r.get_or_add_rPr().append(ea)
        slide.notes_slide.notes_text_frame.text = f"第 {i} 页｜{item.title}\n建议用时：{item.minutes} 分钟\n\n{item.notes}"
    path = OUT / "谁在说话_六年级科技课.pptx"
    p.save(path)
    return path


def pdf(slides):
    pdf_font = "ClassroomChinese"
    try:
        pdfmetrics.registerFont(TTFont(pdf_font, FONT_PATH))
    except Exception:
        pdf_font = "STSong-Light"
        pdfmetrics.registerFont(UnicodeCIDFont(pdf_font))
    c = canvas.Canvas(str(OUT / "谁在说话_六年级科技课.pdf"), pagesize=(960, 540))
    c.setTitle("谁在说话？——六年级科技课（静态预览）")
    for slide in slides:
        c.setFillColor("#" + BG)
        c.rect(0, 0, 960, 540, fill=1, stroke=0)
        for e in slide.elements:
            x, y = e["x"] * 0.6, 540 - e["y"] * 0.6
            if e["kind"] in ("box", "oval"):
                w, h = e["w"] * 0.6, e["h"] * 0.6
                c.setFillColor("#" + e["fill"])
                if e["kind"] == "oval":
                    c.ellipse(x, y - h, x + w, y, fill=1, stroke=0)
                else:
                    c.roundRect(x, y - h, w, h, e["radius"] * 0.6, fill=1, stroke=0)
            elif e["kind"] == "line":
                c.setStrokeColor("#" + e["color"])
                c.setLineWidth(e["weight"] * 0.6)
                c.line(x, y, e["x2"] * 0.6, 540 - e["y2"] * 0.6)
            else:
                c.setFont(pdf_font, e["size"] * 0.6)
                c.setFillColor("#" + e["color"])
                for i, line in enumerate(e["text"].split("\n")):
                    yy = y - e["size"] * 0.6 - i * e["size"] * e["leading"] * 0.6
                    if e["align"] == "center":
                        c.drawCentredString(x + e["w"] * 0.3, yy, line)
                    elif e["align"] == "right":
                        c.drawRightString(x + e["w"] * 0.6, yy, line)
                    else:
                        c.drawString(x, yy, line)
        c.showPage()
    c.save()


def preview(slides):
    previews = []
    for slide in slides:
        im = Image.new("RGB", (W, H), "#" + BG)
        draw = ImageDraw.Draw(im)
        for e in slide.elements:
            x, y = e["x"], e["y"]
            if e["kind"] in ("box", "oval"):
                bounds = (x, y, x + e["w"], y + e["h"])
                if e["kind"] == "oval":
                    draw.ellipse(bounds, fill="#" + e["fill"])
                else:
                    draw.rounded_rectangle(bounds, e["radius"], fill="#" + e["fill"])
            elif e["kind"] == "line":
                draw.line((x, y, e["x2"], e["y2"]), fill="#" + e["color"], width=max(1, round(e["weight"])))
            else:
                f = font(e["size"])
                for i, line in enumerate(e["text"].split("\n")):
                    xx = x
                    width = draw.textlength(line, font=f)
                    if e["align"] == "center":
                        xx += (e["w"] - width) / 2
                    elif e["align"] == "right":
                        xx += e["w"] - width
                    bounds = draw.textbbox((0, 0), line, font=f)
                    draw.text((xx, y + i * e["size"] * e["leading"] - bounds[1] + e["size"] * 0.1),
                              line, font=f, fill="#" + e["color"])
        previews.append(im)
    cols, tw, th, gap = 4, 480, 270, 20
    rows = math.ceil(len(previews) / cols)
    sheet = Image.new("RGB", (cols * (tw + gap) + gap, rows * (th + gap) + gap), "#0B1422")
    for i, im in enumerate(previews):
        sheet.paste(im.resize((tw, th), getattr(Image, "Resampling", Image).LANCZOS), (gap + (i % cols) * (tw + gap), gap + (i // cols) * (th + gap)))
    sheet.save(OUT / "课件预览.png")
    previews[0].save(OUT / "封面.png")
    # 保存一页便于检查课堂中信息量最多的情境题。
    previews[11].save(OUT / "情境题预览.png")


def notes(slides):
    lines = ["谁在说话？——六年级科技展示课\n40 分钟 · 苹果芯片 Mac\n",
             "使用说明：PPT 的音视频位置尚未填入；先按 README.md 制作本人素材，再插入。",
             "A/B 固定顺序为 A 原声、B 合成；竞猜后立即揭晓。第 18–20 页为教师附录。\n"]
    elapsed = 0
    for i, slide in enumerate(slides, 1):
        when = f"{elapsed}–{elapsed + slide.minutes} 分钟" if slide.minutes else "教师附录，不计入课堂时间"
        lines.extend(["=" * 52, f"第 {i} 页｜{slide.title}｜{when}", slide.notes, ""])
        elapsed += slide.minutes
    assert elapsed == 40, elapsed
    (OUT / "逐页讲稿.txt").write_text("\n".join(lines), encoding="utf-8")


def main():
    OUT.mkdir(exist_ok=True)
    compose()
    ppt(SLIDES)
    pdf(SLIDES)
    preview(SLIDES)
    notes(SLIDES)
    print(f"已生成 {len(SLIDES)} 页；课堂部分共 {sum(s.minutes for s in SLIDES)} 分钟。输出：{OUT}")


if __name__ == "__main__":
    main()
