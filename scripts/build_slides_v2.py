#!/usr/bin/env python3
"""《谁在替我说话》：技术展示版，原生可编辑 PPT、PDF 与实际 PDF 页面预览。"""
from __future__ import annotations

from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
import math
import os

import numpy as np
from PIL import Image, ImageFont
import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "课件" / "重制版"
ART = ROOT / "assets" / "design"
PORTRAIT = ART / "fictional-portrait.png"
TEACHER = ROOT / "assets/input/teacher.jpg"
PHOTO = TEACHER if TEACHER.is_file() else PORTRAIT
PAPER, INK, RED = "F2F0E9", "18201E", "C83D31"
MID, LINE, SOFT = "626A65", "D1D4CB", "E3E5DC"
DARK, LIGHT, MINT = "151D1C", "EDEEE6", "B8EDC7"
VOICE, MOTION = "266C57", "B95332"
REG = ROOT / "assets/fonts/NotoSansSC-Regular.ttf"
BOLD = ROOT / "assets/fonts/NotoSansSC-Bold.ttf"
MONO = ROOT / "assets/fonts/JetBrainsMono-Regular.ttf"
W, H = 1600, 900

SOURCES = {
    "voice": "https://github.com/SWivid/F5-TTS",
    "face": "https://github.com/OpenTalker/SadTalker",
    "motion": "https://github.com/KlingAIResearch/LivePortrait",
    "case": "https://www.info.gov.hk/gia/general/202406/26/P2024062600193.htm",
    "provenance": "https://spec.c2pa.org/specifications/specifications/2.2/explainer/Explainer.html",
}


@lru_cache(None)
def measure_font(size, weight="regular"):
    path = MONO if weight == "mono" else BOLD if weight == "bold" else REG
    return ImageFont.truetype(str(path), int(size))


def wrap(text, width, size, weight):
    f = measure_font(size, weight)
    lines = []
    for paragraph in text.split("\n"):
        current = ""
        for ch in paragraph:
            if current and f.getlength(current + ch) > width:
                lines.append(current)
                current = ch
            else:
                current += ch
        lines.append(current)
    return lines


@dataclass
class Slide:
    title: str
    minutes: int
    notes: str
    section: str = ""
    bg: str = PAPER
    elements: list = field(default_factory=list)

    @property
    def fg(self):
        return LIGHT if self.bg == DARK else INK

    @property
    def secondary(self):
        return "AAB7AD" if self.bg == DARK else MID

    def rect(self, x, y, w, h, fill, stroke=None, sw=1, radius=0):
        self.elements.append(dict(k="rect", x=x, y=y, w=w, h=h, fill=fill,
                                  stroke=stroke, sw=sw, radius=radius))

    def circle(self, cx, cy, r, fill, stroke=None, sw=1):
        self.elements.append(dict(k="oval", x=cx-r, y=cy-r, w=2*r, h=2*r,
                                  fill=fill, stroke=stroke, sw=sw))

    def line(self, x1, y1, x2, y2, color=LINE, sw=2, dash=False):
        self.elements.append(dict(k="line", pts=[(x1, y1), (x2, y2)], color=color, sw=sw, dash=dash))

    def path(self, points, color=INK, sw=2, closed=False, fill=None):
        self.elements.append(dict(k="path", pts=points, color=color, sw=sw, closed=closed, fill=fill))

    def text(self, x, y, w, h, text, size=40, color=None, weight="regular", align="left", leading=1.25):
        lines = wrap(text, w, size, weight)
        if len(lines) * size * leading > h + 1:
            raise ValueError(f"需要调整文字高度：{self.title} / {text}")
        self.elements.append(dict(k="text", x=x, y=y, w=w, h=h, text="\n".join(lines),
                                  size=size, color=color or self.fg, weight=weight,
                                  align=align, leading=leading))

    def image(self, path, x, y, w, h, anchor=(.5, .5), movie=None):
        self.elements.append(dict(k="image", path=str(path), x=x, y=y, w=w, h=h,
                                  anchor=anchor, movie=str(movie) if movie and Path(movie).is_file() else None))


SLIDES = []


def page(title, minutes, notes, section, bg=PAPER, heading=True):
    s = Slide(title, minutes, notes, section, bg)
    SLIDES.append(s)
    s.text(76, 43, 1448, 42, section, 22, s.secondary, "regular")
    if heading:
        s.text(76, 108, 1448, 108, title, 62, weight="bold")
    return s


def tag(s, x, y, text, color=RED, width=None):
    w = width or max(92, measure_font(21).getlength(text) + 30)
    s.rect(x, y, w, 37, color)
    s.text(x+15, y+4, w-26, 30, text, 21, LIGHT)


def arrow(s, x1, y, x2, color=INK, sw=3):
    s.line(x1, y, x2, y, color, sw)
    s.line(x2-13, y-9, x2, y, color, sw)
    s.line(x2-13, y+9, x2, y, color, sw)


def play(s, cx, cy, r=43, color=PAPER):
    s.circle(cx, cy, r, color)
    s.path([(cx-10, cy-17), (cx-10, cy+17), (cx+18, cy)], INK, 0, True, INK)


def photo_note(s, x, y, w, text=None, light=False):
    label = text or ("本人肖像" if TEACHER.is_file() else "虚构人物 · AI 生成 · 视觉示意")
    s.text(x, y, w, 36, label, 20, "C0C7BD" if light else MID)


def waveform(s, x, y, w, h, color=VOICE, variant=0):
    # 结构示意，不是假装从真人录音中提取的波形。
    n = 360
    points = []
    for i in range(n):
        t = i/(n-1)
        env = (0.08 + 0.9*math.sin(math.pi*t)**1.1) * (0.35+0.65*math.sin(5*math.pi*t+variant*.2)**2)
        value = math.sin(102*math.pi*t+variant)*.66 + math.sin(174*math.pi*t)*.19
        points.append((x+t*w, y+h/2-value*env*h*.48))
    s.line(x, y+h/2, x+w, y+h/2, LINE if s.bg == PAPER else "3B4942", 1)
    s.path(points, color, 2)


def spectrogram(s, x, y, w, h):
    """计算人工合成谐波信号的短时频谱，明确标为声学表示示意。"""
    sr=8000
    t=np.arange(sr*2)/sr
    f0=145+28*np.sin(2*np.pi*.9*t)
    phase=2*np.pi*np.cumsum(f0)/sr
    envelope=(.10+.9*np.sin(np.pi*t/2)**2)*(.22+.78*np.sin(2*np.pi*1.35*t)**2)
    signal=envelope*sum(np.sin(k*phase)/k for k in range(1,12))
    win=512
    positions=np.linspace(0,len(signal)-win,60).astype(int)
    bins=38
    spectra=[]
    for pos in positions:
        spec=np.abs(np.fft.rfft(signal[pos:pos+win]*np.hanning(win)))
        # 只显示 0–2.5 kHz，聚合相邻频率区间。
        edges=np.linspace(0,int(2500*win/sr),bins+1).astype(int)
        spectra.append([float(np.max(spec[a:max(a+1,b)])) for a,b in zip(edges[:-1],edges[1:])])
    values=np.array(spectra).T
    values=np.clip((20*np.log10(values+1e-5)+23)/60,0,1)
    lo=np.array([27,43,35]);hi=np.array([184,237,199])
    for row in range(bins):
        for col in range(len(positions)):
            v=values[row,col]**1.3
            rgb=np.round(lo*(1-v)+hi*v).astype(int)
            color="".join(f"{a:02X}" for a in rgb)
            s.rect(x+col*w/60,y+(bins-row-1)*h/bins,w/60+.3,h/bins+.3,color)


def face_mesh(s, x, y, w, h, color=MINT):
    # 以虚构肖像为参照的结构图，非检测器/模型实际输出。
    groups = [
        [(.27,.40),(.31,.365),(.37,.36),(.42,.39),(.37,.415),(.31,.416)],
        [(.57,.39),(.62,.36),(.69,.365),(.735,.40),(.69,.417),(.62,.413)],
        [(.5,.4),(.48,.47),(.46,.55),(.50,.565),(.55,.55)],
        [(.37,.65),(.43,.625),(.5,.62),(.565,.633),(.62,.652),(.56,.677),(.50,.686),(.43,.675)],
        [(.22,.39),(.23,.49),(.25,.59),(.30,.69),(.39,.77),(.5,.81),(.61,.77),(.70,.69),(.75,.59),(.78,.49),(.78,.39)],
    ]
    for group in groups:
        pts=[(x+a*w,y+b*h) for a,b in group]
        s.path(pts,color,1.7)
        for px,py in pts:
            s.circle(px,py,2.4,color)
    for a,b in [((.31,.416),(.46,.55)),((.69,.417),(.55,.55)),((.46,.55),(.37,.65)),
                ((.55,.55),(.62,.652)),((.5,.565),(.5,.686)),((.37,.65),(.39,.77)),
                ((.62,.652),(.61,.77)),((.23,.49),(.37,.65)),((.78,.49),(.62,.652))]:
        s.line(x+a[0]*w,y+a[1]*h,x+b[0]*w,y+b[1]*h,color,1)


def build_content():
    s=page("谁在替我说话？",1,
        "开场只提出一个问题：如果声音和脸都能生成，我们凭什么判断一个人真的表达过？\n"
        "今天会展示本人授权制作的数字分身，部分材料明确标注为合成。先看真实效果，再拆机制。\n"
        "封面虚构肖像仅用于视觉示意，不代表讲师，不是假装已经完成了真人克隆。", "AI 安全 · 声音克隆 / 人脸合成",heading=False)
    s.rect(941,0,659,900,"D3D4CC")
    s.image(PHOTO,941,0,659,900,anchor=(.5,.45))
    s.rect(941,541,659,9,RED)
    s.rect(979,528,282,52,RED)
    s.text(996,535,252,40,"SYNTHETIC IDENTITY",22,LIGHT,"mono")
    s.text(76,192,850,344,"谁在替我\n说话？",130,INK,"bold",leading=1.15)
    s.line(82,603,240,603,RED,7)
    s.text(76,639,790,118,"声音克隆、人脸合成\n与身份验证",43,INK)
    s.text(76,790,800,44,"科技报告 · 六年级 · 40 分钟",25,MID)
    photo_note(s,969,846,600,light=True)

    s=page("一个我在讲台上，另一个我在屏幕里。",3,
        "主展示：播放本人数字分身开场片段，建议 10–15 秒，随后可重放一次。\n"
        "建议台词：‘你现在看到的是我的数字分身。这段话没有被我录下来，而是由人工智能生成的。’\n"
        "保持本人在讲台上，形成直观对照。问：画面、声音、台词，这三部分里哪些来自我？哪些由模型生成？\n"
        "本页会在文件存在时嵌入 assets/output/02_talking_ai.mp4；当前若没有该文件，显示的是明确标注的视觉示意海报。\n"
        "PowerPoint 中单击视频播放；PDF 是静态海报。", "01 / 先看效果",DARK)
    s.image(PHOTO,76,248,614,540,anchor=(.5,.38),movie=ROOT/"assets/output/02_talking_ai.mp4")
    has_video=(ROOT/"assets/output/02_talking_ai.mp4").is_file()
    if not has_video:
        play(s,383,518,48)
    tag(s,100,270,"AI 合成 · 本人授权演示" if has_video else "演示素材位 · 待加入本人视频",width=370)
    s.text(760,300,748,210,"“这段话，\n我没有录过。”",74,LIGHT,"bold")
    s.text(766,567,716,120,"但它可以拥有我的声音，\n以及一张看起来像我的脸。",40,"BBC6BA")
    s.text(766,733,720,51,"观察：声音 / 外观 / 口型 / 台词",27,MINT)
    photo_note(s,76,803,1000,"已嵌入本人演示视频" if (ROOT/"assets/output/02_talking_ai.mp4").is_file()
               else "媒体位：待加入本人合成视频；当前肖像仅为 AI 生成视觉示意",True)

    s=page("把一个“像我的人”拆开",1,
        "用刚才片段来回答，不做猜中与猜错的计分。\n"
        "三项来源：声音身份线索来自参考录音；外观来自照片/视频；台词由作者输入。运动则由音频驱动或从动作视频迁移。\n"
        "这一页强调可分离、可重组，并不声称所有系统都能完全独立控制每项特征。", "01 / 三种可组合的信息")
    s.image(PHOTO,76,246,372,552,anchor=(.5,.4))
    photo_note(s,76,800,600)
    items=[("01","声音","听起来像谁",VOICE),("02","外观","看起来像谁",INK),("03","内容","到底说什么",RED)]
    for i,(n,title,desc,c) in enumerate(items):
        yy=270+i*177
        s.line(530,yy+143,1519,yy+143,LINE,1)
        s.text(540,yy,100,74,n,40,c,"mono")
        s.text(674,yy-3,246,92,title,64,c,"bold")
        s.text(1035,yy+14,480,75,desc,40)

    s=page("先拿掉画面，只听声音。",2,
        "对照听音：A 插入 assets/input/teacher_real.wav；B 插入 assets/output/01_voice_ai.wav。均读 comparison_text.txt，音量接近。\n"
        "提问：你依据的是音色、节奏、呼吸，还是已经知道的上下文？允许判断不确定。\n"
        "A 原声、B 合成。竞猜后在当前页口头揭晓，不延迟制造误会。波形只是结构示意，不是两段录音的测量图。\n"
        "模板没有捏造音频；在 PowerPoint/Keynote 中将两段实际音频分别插入 A/B 区域。", "02 / 声音克隆")
    for i,(label,color) in enumerate([("A",INK),("B",RED)]):
        yy=260+i*230
        s.text(80,yy,130,150,label,112,color,"mono")
        waveform(s,260,yy+12,1060,126,color,variant=i)
        play(s,1467,yy+77,40,SOFT)
    s.line(76,720,1524,720,LINE,1)
    s.text(76,752,1448,63,"相同台词，让比较集中在声音本身。",40)
    s.text(260,681,1000,33,"示意波形 · 待插入本人原声 / 合成声",20,MID)

    s=page("声音不是把旧录音剪成新句子。",3,
        "展示主线的输入输出：参考录音、录音逐字稿与新台词，经过 F5-TTS，生成新声音。\n"
        "建议先播放参考录音，再显示一段完全不同的新台词。模型参考声音特点生成新的声学信号，而不是把原录音中的每个字剪出来拼接。\n"
        "零样本指不针对这个说话人专门微调，不是从零训练模型。模型已经从大量语音/文字数据中学习过关系。\n"
        "来源："+SOURCES["voice"], "02 / 内容与声音身份分开")
    s.text(76,270,650,55,"参考录音 + 对应文字",33,MID)
    waveform(s,76,348,620,139,VOICE)
    s.text(76,539,650,52,"想让分身说的新台词",33,MID)
    s.text(76,617,650,128,"“这段话，我没有录过。”",43,RED,"bold")
    arrow(s,745,472,896,INK)
    s.rect(944,269,580,415,INK)
    s.text(986,305,480,50,"F5-TTS / INFERENCE",26,MINT,"mono")
    s.text(986,404,478,130,"新的声音",67,LIGHT,"bold")
    waveform(s,987,565,485,70,MINT,2)
    s.text(76,775,1448,55,"参考样本提供声音线索，新台词决定要说的内容。",36)

    s=page("几秒钟样本，接上一个已经学过的模型。",2,
        "先区分训练与推理。左侧是模型此前的学习，右侧是本次生成。我们本次提供的短录音只是条件信息，不是整个模型的训练数据。\n"
        "简要解释：声音可以表示为波形，也可以用时频表示帮助模型处理；模型学习文字与声学表现的关系，再结合参考语音生成。\n"
        "频谱图是结构示意，不是真人录音的分析结果；无需讲网络层数或优化公式。\n"
        "来源：F5-TTS 官方代码、推理说明。", "02 / 训练 ≠ 本次推理",DARK)
    s.text(76,260,610,70,"模型此前学过",44,MINT,"bold")
    s.text(76,350,610,60,"文字与发音、声学变化的关系",31,LIGHT)
    spectrogram(s,76,430,610,190)
    s.text(76,637,610,75,"横轴：时间　纵轴：频率\n颜色：声音强弱",25,"BBC6BA")
    s.text(76,734,610,40,"人工合成信号的频谱示意 · 非真人录音",20,"8FA794")
    s.line(743,266,743,781,"435149",1)
    s.text(810,260,714,70,"这一次，再提供",44,MINT,"bold")
    for i,(a,b) in enumerate([("音色线索","参考声音"),("内容线索","新台词")]):
        yy=372+i*128
        s.text(810,yy,250,55,a,30,"BBC6BA")
        s.text(1120,yy-10,400,80,b,48,LIGHT,"bold")
    s.rect(810,661,714,112,MINT)
    s.text(840,685,650,67,"→ 生成新语音",47,INK,"bold")

    s=page("再给这段声音一张脸。",2,
        "再次播放照片说话片段，重点看嘴部、眼部表情、头部姿态的变化。\n"
        "主线使用 SadTalker：单张照片给外观信息，音频驱动模型预测面部运动参数，再生成逐帧画面。\n"
        "这里不把人像动画说成狭义换脸。右图是虚构人物外观参照，结构线只是原生示意图。\n"
        "来源："+SOURCES["face"], "03 / 音频驱动人像动画")
    s.image(PHOTO,1006,242,518,562,anchor=(.5,.4))
    face_mesh(s,1006,242,518,562,RED)
    s.text(76,263,848,180,"照片提供外观，\n声音驱动变化。",61,INK,"bold")
    waveform(s,76,468,820,108,VOICE)
    arrow(s,775,640,954,RED)
    s.text(76,644,684,109,"嘴部开合 · 表情 · 头部姿态",38)
    s.text(1011,809,513,32,"面部结构示意 · 非模型检测输出",18,MID)

    s=page("中间发生了什么？",3,
        "按从左到右的顺序，解释声音如何变成会说话的画面。\n"
        "第一步，模型分析音频中的发音与时间变化；第二步，预测面部表情和头部姿态等运动参数；第三步，结合照片生成连续帧。\n"
        "三个小图是运动与连续帧的概念示意，不能当作模型真实可视化。实际 SadTalker 使用三维运动系数；精确网络结构留给教师参考。\n"
        "关键结论：视频里的变化可以被生成；播放的流畅性不是现场拍摄的证明。\n来源：SadTalker 官方论文与仓库。", "03 / 声音 → 运动 → 连续画面")
    columns=[76,581,1086]
    for i,(num,title) in enumerate([("01","分析音频"),("02","预测运动"),("03","合成画面")]):
        xx=columns[i]
        s.text(xx,256,430,60,num,35,RED,"mono")
        s.text(xx,320,438,88,title,50,INK,"bold")
    waveform(s,76,475,420,122,VOICE)
    for j,(label,length) in enumerate([("口型",268),("表情",186),("姿态",225)]):
        yy=460+j*81
        s.text(581,yy,96,54,label,31,MID)
        s.rect(703,yy+8,302,22,SOFT)
        s.rect(703,yy+8,length,22,VOICE if j<2 else RED)
    for j in range(3):
        xx=1086+j*112
        s.rect(xx,458,100,162,SOFT)
        s.circle(xx+50,511,29,INK)
        s.rect(xx+26,543,48,47,INK)
        s.line(xx+39,522,xx+61,522, MINT,2+j*2)
    arrow(s,503,526,554,INK)
    arrow(s,1017,526,1063,INK)
    s.line(76,699,1524,699,LINE,1)
    s.text(76,737,1448,75,"连续播放这些画面，就形成了“正在说话”的视频。",42)
    s.text(76,810,1448,30,"运动参数与帧图均为结构示意，无实测数值含义。",18,MID)

    s=page("“换脸”，其实包含不同的事。",2,
        "按输入与变化部分区分，避免把所有人脸生成统称为一个技术。\n"
        "换脸：在已有视频中替换人物身份外观；照片说话：让静态肖像随音频产生运动；表情迁移：把驱动视频的运动转到目标肖像。\n"
        "本课主展示是声音克隆 + 照片说话，可选扩展是 LivePortrait 表情迁移。LivePortrait 不会直接根据一段新音频自动做对应口型。\n"
        "来源：SadTalker、LivePortrait 官方仓库。", "03 / 三类技术的边界")
    labels=[("换脸","已有视频 + 另一张脸","人物外观被替换",INK),
            ("照片说话","肖像 + 音频","静态照片产生说话动作",VOICE),
            ("表情迁移","肖像 + 驱动视频","把动作带到另一张脸上",RED)]
    for i,(a,b,c,col) in enumerate(labels):
        yy=269+i*177
        s.line(76,yy+144,1524,yy+144,LINE,1)
        s.text(76,yy,325,88,a,54,col,"bold")
        s.text(437,yy+9,620,62,b,37)
        s.text(1093,yy+12,431,110,c,32,MID)
    s.text(76,802,1448,34,"主展示：F5-TTS + SadTalker；扩展：LivePortrait。",23,MID)

    s=page("这次展示，用到的是这条制作链。",2,
        "把上一段技术解释落回你实际使用的项目，展示输入与产物。\n"
        "本人参考录音与准确逐字稿、新台词，经 F5-TTS 生成声音；本人照片与生成音频，经 SadTalker 生成视频；最后加入明确合成标识。\n"
        "在 Mac 上，声音可尝试 MPS；SadTalker 官方当前入口选择 CUDA 或 CPU，因此本项目 Mac 视频方案走 CPU，安排在课前。\n"
        "不要承诺固定秒数或实时视频。现场改台词演示可作为已实测后的附加环节。", "04 / 从输入到成片")
    y1,y2=298,552
    s.text(76,242,700,46,"声音路径",29,VOICE,"bold")
    s.text(76,y1,510,118,"参考录音 + 逐字稿\n+ 新台词",38)
    s.rect(652,y1-5,361,124,INK)
    s.text(682,y1+24,301,64,"F5-TTS",43,MINT,"mono",align="center")
    arrow(s,547,y1+58,618,INK)
    arrow(s,1043,y1+58,1115,INK)
    s.text(1160,y1+22,360,77,"合成声音",46,VOICE,"bold")
    s.line(76,477,1524,477,LINE,1)
    s.text(76,496,700,46,"画面路径",29,RED,"bold")
    s.text(76,y2+16,510,99,"本人照片 + 合成声音",38)
    s.rect(652,y2,361,124,INK)
    s.text(682,y2+28,301,64,"SadTalker",38,LIGHT,"mono",align="center")
    arrow(s,547,y2+62,618,INK)
    arrow(s,1043,y2+62,1115,INK)
    s.text(1160,y2+28,360,77,"说话视频",46,RED,"bold")
    s.text(76,764,1448,60,"先在课前生成视频，课堂展示输入、结果和关键变化。",35,MID)

    s=page("如果它还能现场回答问题呢？",2,
        "先问：一段预录视频与实时数字人，增加了哪些环节？\n"
        "讲解实时链：接收语音→识别文字→生成回答→合成声音→驱动人脸动画。此图是系统组成，不是声称当前这套项目已经实现低延迟实时数字人。\n"
        "追问：能接话、能转头，可以排除某些简单重放，但能否单独证明身份和授权？不能把互动能力当绝对身份证明。\n"
        "如展示实时系统，应另外说明所用软件、设备与延迟；本课默认视频预生成。", "04 / 从预录片段到实时交互",DARK)
    s.text(76,257,1448,180,"能接话，增加了互动；\n身份仍然需要验证。",65,LIGHT,"bold")
    stages=["听见问题","理解与回答","合成声音","生成画面"]
    for i,label in enumerate(stages):
        xx=76+i*371
        s.circle(xx+24,553,8,MINT)
        if i<3: arrow(s,xx+47,553,xx+337,"5C7463",2)
        s.text(xx,601,340,80,label,36,LIGHT,"bold")
    s.text(76,757,1448,66,"系统组成示意 · 本课演示主线为课前生成片段",29,"AAB7AD")

    s=page("视频会议里的人，也可能来自生成。",2,
        "真实案例，按香港特区政府 2024-06-26 立法会书面答复概述：2024 年 1 月底警方接报，涉案人员收到假冒英国总公司首席财务官的钓鱼消息，被邀加入利用深伪预录影像构成的多人视频会议，随后授权向多个本地账户转账，损失达数亿港元。\n"
        "官方后续不同时间点金额口径有更新；课堂不追逐金额，重点讲信任链如何被利用。\n"
        "不能把该案描述成已证实的实时换脸会议。图中只是事件结构，没有冒充案件实际截图。\n"
        "提问：如果会议里看似很多人都同意了，为什么仍可能缺少独立证据？\n来源："+SOURCES["case"], "05 / 真实案例")
    s.text(76,260,590,105,"2024 / HONG KONG",35,RED,"mono")
    s.text(76,370,610,198,"多人“在场”，\n来源可能只有一方。",56,INK,"bold")
    s.text(76,680,610,119,"问题出在：外观、声音与权威感，\n被当成了身份和授权的证据。",30,MID)
    xx=826
    for i,(a,b) in enumerate([("钓鱼消息","假冒公司高层"),("预录深伪会议","制造熟人和多人参与的情境"),("授权转账","把指令转成真实行动")]):
        yy=261+i*176
        s.circle(xx,yy+35,16,RED if i==2 else INK)
        if i<2:s.line(xx,yy+55,xx,yy+165,LINE,2)
        s.text(xx+48,yy,640,75,a,46,RED if i==2 else INK,"bold")
        s.text(xx+48,yy+82,640,71,b,30,MID)
    s.text(76,814,1448,31,"案例结构示意；来源：香港特区政府立法会书面答复，2024-06-26。",18,MID)

    s=page("眨眼、转头、说口令，能证明什么？",3,
        "请学生评价三种建议，不急着给出万能答案：要求对方眨眼、说随机句子、换已知可信渠道联系。\n"
        "眨眼/转头：可以增加运动线索，但动画模型也能生成这些动作。随机回应：可能提高对预录片段的辨别力，但实时生成系统同样可能互动。\n"
        "已知可信渠道确认，可以增加独立证据，但安全仍取决于渠道是否受控和流程是否执行。\n"
        "这里不否定活体检测的全部价值；讲的是普通人临时指定动作，不等于完整可靠的身份认证。", "05 / 交互能力与身份认证")
    rows=[("眨眼、转头","证明有动作？","动作也可以生成"),
          ("说一句随机的话","排除简单重放？","实时系统也可能回应"),
          ("通过已知可信渠道核实","增加独立证据？","还要确认渠道与授权流程")]
    for i,(a,b,c) in enumerate(rows):
        yy=272+i*166
        s.line(76,yy+136,1524,yy+136,LINE,1)
        s.text(76,yy,526,93,a,39,INK,"bold")
        s.text(650,yy+4,406,95,b,32,VOICE)
        s.text(1128,yy+5,396,110,c,30,MID)
    s.text(76,793,1448,43,"一项动作检查，只回答其中一部分问题。",29,RED)

    s=page("“真不真”，至少有三层意思。",2,
        "区分三个问题：制作方式是否为真实拍摄；是否代表本人授权表达；内容是否符合事实。\n"
        "举例：本人授权的合成科普，可以没有原始拍摄却有真实授权；真实拍摄的人也可能讲错事实；未经同意的合成视频，即使其中某句事实正确，仍可能侵犯身份表达。\n"
        "学生不必背术语，能说清自己究竟在判断哪件事即可。", "06 / 把问题拆准确")
    for i,(num,title,question) in enumerate([
        ("01","制作方式","是拍摄，还是生成？"),
        ("02","身份与授权","是否由本人认可？"),
        ("03","内容事实","说的事情是否属实？"),
    ]):
        xx=76+i*502
        s.text(xx,276,430,98,num,74,RED,"mono")
        s.line(xx,405,xx+432,405,LINE,2)
        s.text(xx,448,436,110,title,50,INK,"bold")
        s.text(xx,611,430,122,question,36,MID)
    s.rect(76,776,1448,53,INK)
    s.text(100,783,1400,40,"判断“合成”本身，回答不了所有问题。",27,LIGHT)

    s=page("再让 AI 检测 AI，就能结束争论吗？",2,
        "本页只有概念混淆矩阵，没有虚构准确率、检测分数或实验统计。\n"
        "横轴是实际来源，纵轴是检测器给出的判断。可能把真实拍摄误判为合成，也可能漏掉合成内容。\n"
        "检测效果还与数据、压缩、分辨率、生成模型及使用条件有关；不引用某个演示小样本的高准确率作普遍承诺。\n"
        "关键：检测器是一种证据工具，它不能单独判定本人授权，也不直接判定视频里的说法是否真实。", "06 / 检测是证据之一")
    s.text(76,274,500,195,"它也会\n误报和漏报。",64,INK,"bold")
    s.text(76,592,500,165,"一个检测结果，\n还需要结合来源、上下文\n和其他独立证据。",33,MID)
    gx,gy,cw,ch=806,358,320,180
    s.text(gx,272,cw,60,"实际：拍摄",30,INK,"bold",align="center")
    s.text(gx+cw,272,cw,60,"实际：合成",30,INK,"bold",align="center")
    for i in range(2):
        for j in range(2):
            bad=i!=j
            s.rect(gx+j*cw,gy+i*ch,cw,ch,"F0D9CD" if bad else SOFT, PAPER,4)
            txt=[["正确识别","漏报"],["误报","正确识别"]][i][j]
            s.text(gx+j*cw+18,gy+i*ch+55,cw-36,76,txt,41,RED if bad else VOICE,"bold",align="center")
    s.text(630,406,160,100,"判断为\n拍摄",27,MID,align="right")
    s.text(630,586,160,100,"判断为\n合成",27,MID,align="right")
    s.text(gx,763,670,41,"概念图 · 不包含实测准确率",22,MID)

    s=page("怎样为“是本人”增加证据？",3,
        "把防御讲成系统：第一，能独立核实的联系方式或当面确认；第二，能够验证发布者/签署者的来源记录；第三，重要行动仍通过授权与复核流程。\n"
        "数字签名可帮助确认谁用某个密钥签署了记录及记录是否遭篡改，但不能自动证明镜头里的事情真实，也不能让密钥保管与签发身份问题消失。\n"
        "C2PA 用于内容来源与编辑记录，其验证不等同于内容事实正确或本人授权。这里只介绍用途，不宣称当前演示视频已带 C2PA 凭证。\n"
        "引导学生提出具体渠道，区别‘同一个可疑聊天里再问一次’与获得新的独立证据。\n来源："+SOURCES["provenance"], "06 / 从内容检查走向信任链",DARK)
    steps=[("来源","从哪里来？","核实已知渠道",MINT),
           ("身份","谁在发布或签署？","验证可信来源记录",MINT),
           ("授权","能据此执行什么？","重要行动独立确认", "F3B5A5")]
    for i,(title,question,action,accent) in enumerate(steps):
        yy=264+i*169
        s.text(76,yy,240,91,title,61,accent,"bold")
        s.text(425,yy+3,554,87,question,40,LIGHT)
        s.text(1090,yy+9,434,92,action,33,"BBC6BA")
        s.line(76,yy+139,1524,yy+139,"3B4942",1)
    s.text(76,801,1448,40,"来源凭证帮助验证记录；内容是否属实仍要看证据。",26,"AAB7AD")

    s=page("同样的能力，边界由谁决定？",2,
        "从技术能力转到有意义的用途：本人授权配音、跨语言讲解、角色动画、辅助表达等。这里只讲使用方向，不声称本项目符合医疗产品要求。\n"
        "对照未经同意的冒充、制造未说过的话、伪造认可与交易指令。\n"
        "讨论标准：本人同意什么用途？观众是否被告知？发布者是否借用别人的身份取得不应有的信任？\n"
        "重点放在授权与表达，不把所有生成内容一概等同于虚假信息。", "07 / 技术、表达与责任")
    s.text(76,263,698,100,"本人授权的表达",53,VOICE,"bold")
    s.text(884,263,640,100,"借用身份取得信任",49,RED,"bold")
    s.line(786,263,786,742,LINE,2)
    s.text(76,410,650,215,"配音与跨语言讲解\n获授权的角色创作\n辅助表达",41,INK,leading=1.5)
    s.text(884,410,640,215,"冒充本人发言\n伪造认可与承诺\n让他人据此采取行动",41,INK,leading=1.5)
    s.text(76,779,1448,60,"谁同意了什么用途？观众知道什么？后果由谁承担？",35)

    s=page("我授权分身读自己的讲稿，算“假吗”？",2,
        "保留两分钟开放讨论。这不是只有一个标准答案的选择题。\n"
        "给定明确条件：稿子由本人写或审核，本人授权数字分身朗读，发布时标注 AI 合成。\n"
        "引导区分：不是原始拍摄；声音由模型生成；但表达内容可以确实得到本人认可。学生可以分别讨论‘我说过’、‘我表达过’、‘我授权发布过’之间的区别。\n"
        "追问：如果原样的视频被去掉标注、换一个账号转发，读者的判断会发生什么变化？不要把去标注说成改变了画面事实，而是上下文与来源信息变化。", "07 / 留一个值得争论的问题",heading=False)
    s.text(76,178,1448,274,"我授权数字分身，\n读出我认可的讲稿。",89,INK,"bold",leading=1.2)
    s.text(76,514,1448,158,"你会把它叫作“假的”吗？",76,RED,"bold")
    s.line(76,738,1524,738,LINE,1)
    s.text(76,776,1448,56,"讨论依据：制作方式 / 本人授权 / 内容事实",33,MID)

    s=page("声音和脸可以生成。信任需要更多证据。",1,
        "收束全课：今天见到的不是剪贴魔术，而是根据参考声音、台词与照片生成新内容的模型能力。\n"
        "回扣三件事：知道它怎样生成；理解检测的边界；在身份和重要授权上寻找独立证据。\n"
        "最后一句：当声音与外观不再足以单独证明身份，我们需要更清楚地知道信息从哪里来、谁认可它、我们据此要做什么。", "AI 安全 · 生成 / 检测 / 认证",DARK,heading=False)
    s.text(76,184,1448,310,"声音和脸\n可以生成。",107,LIGHT,"bold",leading=1.13)
    s.text(76,544,1448,142,"信任需要更多证据。",83,MINT,"bold")
    s.text(76,772,1448,60,"从哪里来 · 谁认可 · 据此做什么",34,"AAB7AD")

    s=page("教师附录：四份素材，完成主展示",0,
        "课堂播放素材尚需本人录制与生成，本课件不假装已经包含真人克隆结果。\n"
        "teacher_reference.wav：本人 8–10 秒参考声音，完整短句、无背景音乐，总长少于 12 秒。teacher_reference.txt 必须与实际录音逐字对应。\n"
        "teacher_real.wav：本人读 comparison_text.txt 的原声；teacher.jpg：正面单人照。\n"
        "先生成 A/B 用的 01_voice_ai.wav，再用未由本人录制的 demo_text.txt 生成 00_opening_voice_ai.wav，后者用于开场视频。\n"
        "按 README.md 的 voice/talking 入口生成，视频文件存在时本生成器可自动嵌入第 2 页。A/B 音频在第 4 页手动嵌入，顺序固定 A 原声、B 合成。\n"
        "可选 teacher_driving.mp4 用于 LivePortrait；本版主线不再为可选动作迁移单独占一整个演示段落。", "教师附录 / 不计入 40 分钟")
    rows=[("本人参考录音 + 逐字稿","teacher_reference.wav / .txt","声音线索"),
          ("比较台词 / 开场台词","comparison_text.txt / demo_text.txt","A/B 同句对照；开场另用未录过的新台词"),("本人原声对照","teacher_real.wav","与 A/B 合成声音读同一段话"),
          ("本人肖像","teacher.jpg","正面、单人、嘴自然闭合")]
    for i,(a,b,c) in enumerate(rows):
        yy=261+i*136
        s.line(76,yy+112,1524,yy+112,LINE,1)
        s.text(76,yy,600,68,a,36,INK,"bold")
        s.text(720,yy+3,800,55,b,27,VOICE,"mono")
        s.text(720,yy+57,800,54,c,27,MID)

    s=page("教师附录：演示与设备边界",0,
        "主线代码、Mac 分环境安装命令与许可证说明见根目录 README.md。源码版本见 projects/SOURCES.json。\n"
        "F5-TTS：官方有苹果芯片支持，代码支持 MPS；速度与成功率仍以你的 Mac 和素材实测为准。\n"
        "SadTalker：当前官方入口在 Mac 选择 CPU，本课前生成视频。\n"
        "LivePortrait：苹果芯片人像模式可作扩展；输入动作视频，不是直接输入新音频生成口型。\n"
        "现场可改一小段台词，但视频使用预生成文件。首次权重下载和完整运行都在课前完成。\n"
        "授权：F5-TTS 官方预训练权重为非商业许可；LivePortrait 所附 InsightFace 模型限制非商业研究；依赖各有条款，详见原仓库。", "教师附录 / 实现与课堂节奏")
    rows=[("F5-TTS","参考声样 + 台词 → 声音","Mac / MPS；先完成本人素材实测"),
          ("SadTalker","肖像 + 音频 → 说话视频","Mac / CPU；视频课前生成"),
          ("LivePortrait","肖像 + 驱动视频 → 动画","扩展项目；不自动按新音频对口型")]
    for i,(a,b,c) in enumerate(rows):
        yy=271+i*174
        s.text(76,yy,450,86,a,46,VOICE if i==0 else INK,"mono")
        s.text(624,yy,900,70,b,38,INK,"bold")
        s.text(624,yy+82,900,68,c,29,MID)
        s.line(76,yy+145,1524,yy+145,LINE,1)
    s.text(76,809,1448,32,"当前真人素材与 Mac 推理尚未完成；界面图解不冒充模型实测。",20,RED)

    s=page("教师附录：来源与图解说明",0,
        "项目接口与能力依据本目录已保存源码及官方说明。参考日期：2026-10-19。\n"
        +"\n".join(f"{key}: {url}" for key,url in SOURCES.items())+
        "\n\n肖像为 imagegen 生成的虚构成年人，生成提示词与设计说明见 assets/design/设计说明.md。波形、面部结构、运动参数和检测矩阵为概念示意；没有伪造实测分数。\n"
        "PDF 与预览从同一版面数据生成，PPT 文本/图形可编辑；当前环境未实际运行 Mac PowerPoint 渲染。", "教师附录 / 官方项目与一手资料")
    refs=[("F5-TTS · 声音生成",SOURCES["voice"]),("SadTalker / LivePortrait · 人像动画",SOURCES["face"]+"\n"+SOURCES["motion"]),
          ("深伪视频会议案例 · 香港特区政府",SOURCES["case"]),("来源记录与内容凭证 · C2PA",SOURCES["provenance"])]
    for i,(a,b) in enumerate(refs):
        yy=252+i*145
        s.text(76,yy,1448,55,a,30,INK,"bold")
        s.text(76,yy+59,1448,76,b,20,VOICE,"mono")

    # 每页只保留细小的定位信息，把注意力留给演示与论证。
    total=len(SLIDES)
    for i,s in enumerate(SLIDES,1):
        if i==1: continue
        col="697B6D" if s.bg==DARK else LINE
        s.line(76,858,1524,858,col,1)
        s.text(76,871,1200,25,"谁在替我说话 / 声音克隆与人脸合成" if s.minutes else "教师资料",15,s.secondary)
        s.text(1390,867,134,30,f"{i:02d}/{total:02d}",18,s.secondary,"mono",align="right")
    assert sum(s.minutes for s in SLIDES)==40


def crop_geometry(path, x,y,w,h,anchor):
    with Image.open(path) as im:
        iw,ih=im.size
    scale=max(w/iw,h/ih)
    dw,dh=iw*scale,ih*scale
    dx=x-(dw-w)*anchor[0]
    dy=y-(dh-h)*anchor[1]
    return dx,dy,dw,dh,(iw,ih)


def pdf(slides):
    pdfmetrics.registerFont(TTFont("Noto",str(REG)))
    pdfmetrics.registerFont(TTFont("NotoBold",str(BOLD)))
    pdfmetrics.registerFont(TTFont("JetBrains",str(MONO)))
    path=OUT/"谁在替我说话.pdf"
    c=canvas.Canvas(str(path),pagesize=(960,540),pageCompression=1)
    c.setTitle("谁在替我说话？——声音克隆、人脸合成与身份验证")
    for slide in slides:
        c.setFillColor("#"+slide.bg); c.rect(0,0,960,540,fill=1,stroke=0)
        for e in slide.elements:
            k=e["k"]
            if k in ["rect","oval"]:
                x,y,w,h=e["x"]*.6,540-(e["y"]+e["h"])*.6,e["w"]*.6,e["h"]*.6
                if e["fill"]:c.setFillColor("#"+e["fill"])
                if e["stroke"]:c.setStrokeColor("#"+e["stroke"]);c.setLineWidth(e["sw"]*.6)
                if k=="oval":c.ellipse(x,y,x+w,y+h,fill=bool(e["fill"]),stroke=bool(e["stroke"]))
                else:c.roundRect(x,y,w,h,e["radius"]*.6,fill=bool(e["fill"]),stroke=bool(e["stroke"]))
            elif k in ["line","path"]:
                c.setStrokeColor("#"+e["color"]); c.setLineWidth(e["sw"]*.6)
                if e.get("fill"):c.setFillColor("#"+e["fill"])
                c.setDash([4,3] if e.get("dash") else [])
                p=c.beginPath()
                for j,(x,y) in enumerate(e["pts"]):
                    (p.moveTo if j==0 else p.lineTo)(x*.6,540-y*.6)
                if e.get("closed"):p.close()
                c.drawPath(p,fill=bool(e.get("fill")),stroke=e["sw"]>0)
                c.setDash([])
            elif k=="image":
                dx,dy,dw,dh,_=crop_geometry(e["path"],e["x"],e["y"],e["w"],e["h"],e["anchor"])
                c.saveState(); p=c.beginPath();p.rect(e["x"]*.6,540-(e["y"]+e["h"])*.6,e["w"]*.6,e["h"]*.6)
                c.clipPath(p,stroke=0)
                c.drawImage(e["path"],dx*.6,540-(dy+dh)*.6,dw*.6,dh*.6,mask="auto")
                c.restoreState()
            elif k=="text":
                name={"regular":"Noto","bold":"NotoBold","mono":"JetBrains"}[e["weight"]]
                c.setFont(name,e["size"]*.6);c.setFillColor("#"+e["color"])
                for j,line in enumerate(e["text"].split("\n")):
                    y=540-(e["y"]+e["size"]*.91+j*e["size"]*e["leading"])*.6
                    if e["align"]=="center": c.drawCentredString((e["x"]+e["w"]/2)*.6,y,line)
                    elif e["align"]=="right":c.drawRightString((e["x"]+e["w"])*.6,y,line)
                    else:c.drawString(e["x"]*.6,y,line)
        c.showPage()
    c.save();return path


def ppt(slides):
    prs=Presentation()
    prs.slide_width,prs.slide_height=Inches(13+1/3),Inches(7.5)
    prs.core_properties.title="谁在替我说话？"
    prs.core_properties.subject="声音克隆、人脸合成与身份验证｜40 分钟科技报告"
    prs.core_properties.author="科技报告"
    inch=lambda a: Inches(a/120)
    rgb=lambda a: RGBColor.from_string(a)
    for i,s in enumerate(slides,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(s.bg)
        for e in s.elements:
            k=e["k"]
            if k in ["rect","oval"]:
                typ=MSO_SHAPE.OVAL if k=="oval" else MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE
                shape=slide.shapes.add_shape(typ,inch(e["x"]),inch(e["y"]),inch(e["w"]),inch(e["h"]))
                if e["fill"]:shape.fill.solid();shape.fill.fore_color.rgb=rgb(e["fill"])
                else:shape.fill.background()
                if e["stroke"]:shape.line.color.rgb=rgb(e["stroke"]);shape.line.width=Pt(e["sw"]*.6)
                else:shape.line.fill.background()
                if k=="rect" and e["radius"]:shape.adjustments[0]=min(.5,e["radius"]/min(e["w"],e["h"]))
            elif k in ["line","path"]:
                pts=e["pts"]
                if k=="path" and (e.get("closed") or len(pts)>2):
                    fb=slide.shapes.build_freeform(start_x=pts[0][0],start_y=pts[0][1],scale=inch(1))
                    fb.add_line_segments(pts[1:],close=e.get("closed",False));shape=fb.convert_to_shape()
                    if e.get("fill"):shape.fill.solid();shape.fill.fore_color.rgb=rgb(e["fill"])
                    else:shape.fill.background()
                else:
                    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,inch(pts[0][0]),inch(pts[0][1]),inch(pts[-1][0]),inch(pts[-1][1]))
                if e["sw"]>0:shape.line.color.rgb=rgb(e["color"]);shape.line.width=Pt(e["sw"]*.6)
                else:shape.line.fill.background()
            elif k=="image":
                if e["movie"]:
                    slide.shapes.add_movie(e["movie"],inch(e["x"]),inch(e["y"]),inch(e["w"]),inch(e["h"]),poster_frame_image=e["path"],mime_type="video/mp4")
                else:
                    _,_,dw,dh,_=crop_geometry(e["path"],e["x"],e["y"],e["w"],e["h"],e["anchor"])
                    pic=slide.shapes.add_picture(e["path"],inch(e["x"]),inch(e["y"]),width=inch(e["w"]),height=inch(e["h"]))
                    pic.crop_left=(dw-e["w"])*e["anchor"][0]/dw
                    pic.crop_right=(dw-e["w"])*(1-e["anchor"][0])/dw
                    pic.crop_top=(dh-e["h"])*e["anchor"][1]/dh
                    pic.crop_bottom=(dh-e["h"])*(1-e["anchor"][1])/dh
            elif k=="text":
                shape=slide.shapes.add_textbox(inch(e["x"]),inch(e["y"]),inch(e["w"]),inch(e["h"]))
                tf=shape.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
                for j,line in enumerate(e["text"].split("\n")):
                    para=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    para.alignment={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}[e["align"]]
                    para.line_spacing=Pt(e["size"]*e["leading"]*.6);para.space_before=para.space_after=Pt(0)
                    run=para.add_run();run.text=line;run.font.name="Menlo" if e["weight"]=="mono" else "PingFang SC"
                    run.font.size=Pt(e["size"]*.6);run.font.bold=e["weight"]=="bold";run.font.color.rgb=rgb(e["color"])
                    ea=OxmlElement("a:ea");ea.set("typeface","PingFang SC");run._r.get_or_add_rPr().append(ea)
        slide.notes_slide.notes_text_frame.text=f"第 {i} 页｜{s.title}\n用时：{s.minutes} 分钟\n\n{s.notes}"
    target=OUT/"谁在替我说话.pptx";prs.save(target);return target


def previews(pdf_path):
    doc=pymupdf.open(pdf_path)
    thumbs=[]
    for i,p in enumerate(doc):
        pix=p.get_pixmap(matrix=pymupdf.Matrix(1,1),alpha=False)
        im=Image.frombytes("RGB",(pix.width,pix.height),pix.samples)
        thumbs.append(im)
        if i in [0,4,7,14,17]:
            pix2=p.get_pixmap(matrix=pymupdf.Matrix(1600/960,1600/960),alpha=False)
            pix2.save(str(OUT/({0:"封面.png",4:"声音原理.png",7:"人脸原理.png",14:"检测边界.png",17:"讨论题.png"}[i])))
    gap=18;tw,th=480,270;cols=4;rows=math.ceil(len(thumbs)/cols)
    sheet=Image.new("RGB",(cols*(tw+gap)+gap,rows*(th+gap)+gap),"#D8DAD2")
    for i,im in enumerate(thumbs):
        sheet.paste(im.resize((tw,th),getattr(Image,"Resampling",Image).LANCZOS),(gap+(i%cols)*(tw+gap),gap+(i//cols)*(th+gap)))
    sheet.save(OUT/"课件总览.png")
    indices=[0,4,7,17]
    key=Image.new("RGB",(2*960+3*gap,2*540+3*gap),"#D8DAD2")
    for i,index in enumerate(indices):key.paste(thumbs[index],(gap+(i%2)*(960+gap),gap+(i//2)*(540+gap)))
    key.save(OUT/"重点页面.png")
    doc.close()


def presenter_notes():
    elapsed=0;lines=["谁在替我说话？——声音克隆、人脸合成与身份验证", "40 分钟技术展示版｜19 页课堂内容 + 3 页教师附录", "",
        "讲授方式：先展示真实效果，再解释机制，最后讨论检测与认证。", "本版肖像为明确标注的 AI 虚构人物。真人音视频尚需本人录制、生成并插入。", ""]
    for i,s in enumerate(SLIDES,1):
        span=f"{elapsed}–{elapsed+s.minutes} 分钟" if s.minutes else "教师附录"
        lines.extend(["="*55,f"{i:02d}｜{s.title}｜{span}",s.notes,""])
        elapsed+=s.minutes
    (OUT/"讲稿与演示安排.txt").write_text("\n".join(lines),encoding="utf-8")


def validate_layout():
    for s in SLIDES:
        for e in s.elements:
            if e["k"] in ["path","line"]:
                assert all(-1<=x<=1601 and -1<=y<=901 for x,y in e["pts"]),(s.title,e)
            else:
                assert 0<=e["x"] and 0<=e["y"] and e["x"]+e["w"]<=1601 and e["y"]+e["h"]<=901,(s.title,e)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    build_content();validate_layout()
    p=ppt(SLIDES);f=pdf(SLIDES);previews(f);presenter_notes()
    print(f"完成：{len(SLIDES)} 页，课堂用时 {sum(s.minutes for s in SLIDES)} 分钟。")
    print(p);print(f)


if __name__=="__main__":main()
